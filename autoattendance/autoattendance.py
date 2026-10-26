import openpyxl

date = input("Date? \n") # input the date of the attendance

# prepare input and output excels
inpath = "input-autoattendance.xlsx" # input Microsoft forms excel
outpath = "output-autoattendance.xlsx" # reference attendance excel to insert the new date into
input = openpyxl.load_workbook(inpath)
output = openpyxl.load_workbook(outpath)
insheet = input.active
outsheet = output["Math Club"]

# basic sorting algorithm - good enough
def bubble_sort(array):
    tries = 1
    while tries:
        tries = 0
        for i in range(len(array) - 1):
            if array[i][0] > array[i+1][0]:
                tries += 1
                swap(array, i, i+1)
            if array[i][0] == array[i+1][0] and array[i][1] and array[i+1][1]:
                if array[i][1] > array[i+1][1]:
                    tries += 1
                    swap(array, i, i+1)
    return array

def swap(array, one, two):
    array[one], array[two] = array[two], array[one]
    return array

# create a new column for the date
outsheet.cell(1, outsheet.max_column+1).value = date

# find existing names in the reference attendance excel
existingnames = []
for row in range(9, outsheet.max_row+1):
    existingnames.append([outsheet.cell(row, 2).value, outsheet.cell(row,1).value])

names = []
for i in range(2, insheet.max_row+1):
    name = insheet.cell(i, 5).value.split(" ") # find the first and last name of the student based on Microsoft form excel

    # if the person is new to math club - insert them into the reference attendance excel in alphabetical order
    if name not in existingnames:
        existingnames.insert(0, name)
        bubble_sort(existingnames)
        new_index = existingnames.index(name)+9
        outsheet.insert_rows(new_index)
        outsheet.cell(new_index, 1).value = name[1]
        outsheet.cell(new_index, 2).value = name[0]

    outsheet.cell(existingnames.index(name)+9, outsheet.max_column).value = "x" # mark them as present

output.save("autoattendance_output.xlsx") # save as separate from the reference attendance excel to avoid overwriting it
